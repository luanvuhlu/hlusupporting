# -*- coding: utf-8 -*- 
from xlrd import open_workbook
import xlrd
from docx import Document
import re
from datetime import datetime

class TimeTableFile:
	subjects=[]
	def __init__(self, file_contents):
		self.file_contents=file_contents
		self.set_subjects()
	def get_subject_time(self, str):
		tmp=str.split("\n")[0].split()
		try:
# 			start=datetime.strptime(tmp[1], '%d/%m/%Y')
# 			end=datetime.strptime(tmp[3], '%d/%m/%Y:')
			start=tmp[1]
			end=tmp[3]
		except:
			return ['', '']
		return [start, end]
	def get_days(self, str):
		tmp=str.split('\n')
		days=[]
		for i in range(1, len(tmp)):
            day = {'week_day': 0, 'start': 0, 'end': 0, 'location': '', 'room': 0}
			d=tmp[i].split()
            day['week_day'] = d[1]
			day['start']=d[3].split(',')[0]
			day['end']=d[3].split(',')[1]
			day['location']=d[7]
			day['room']=re.search('\d+$', d[5]).group(0)
			days.append(day)
		return days
	def set_subjects(self):
		tkb = open_workbook(file_contents=self.file_contents)
		sheet=tkb.sheet_by_index(0)
		found=False
		tmp_subjects=[]
		for r_index in range(sheet.nrows):
			for c_index in range(sheet.ncols):
				value=sheet.cell(r_index, c_index).value
				if value== u'STT':
					row=r_index
					col=c_index
					found=True
					continue
			if found:
				break;
			if value == u'Sinh viên :':
				self.student_name = sheet.cell(r_index, c_index+2).value
				continue
            # if value == u'Mã số :': #Cannot find this cell
            # self.student_code=sheet.cell(r_index, c_index+1).value
            #    		print "Timetable code line 57"+self.student_code
            # continue
			if value== u'Lớp :':
				self.student_class=sheet.cell(r_index, c_index+2).value
				continue
		for r in range(row+1, sheet.nrows):
				if sheet.cell_type(r, col) != xlrd.XL_CELL_NUMBER:
					break;
				subject={'code':'', 'name':'', 'class1':'', 'class2':'', 'time1':'', 'time2': ''}
				if sheet.cell_type(r, col+1) != xlrd.XL_CELL_EMPTY:
					subject['code']=sheet.cell(r, col+1).value
					subject['name']=sheet.cell(r, col+3).value
					class1=sheet.cell(r, col+5).value
					subject['class1']=re.search('\(([a-zA-Z0-9.])+\)$', class1).group(0).split('(')[1].split(')')[0]
					subject['time1']=sheet.cell(r, col+7).value
					if sheet.cell_type(r+1, col+1) == xlrd.XL_CELL_EMPTY:
						subject['time2']=sheet.cell(r+1, col+7).value
						class2=sheet.cell(r+1, col+5).value
						subject['class2']=re.search('\(([a-zA-Z0-9.])+\)$', class2).group(0).split('(')[1].split(')')[0]
					tmp_subjects.append(subject)
		for tmp_sub in tmp_subjects:
			subject={'name':'', 'theory': '', 'seminar':'', 'start':'', 'end':'', 'day_theories':[], 'day_seminars':[]}
			subject['name']=tmp_sub['name']
			subject_time=self.get_subject_time(tmp_sub['time1'])
			subject['start']=subject_time[0]
			subject['end']=subject_time[1]
            if len(tmp_sub['class2']) > 0 and len(tmp_sub['class1']) > len(tmp_sub['class2']):
                subject['seminar']=tmp_sub['class1']
				subject['theory']=tmp_sub['class2']
				day_seminars=self.get_days(tmp_sub['time1'])
				day_theories=self.get_days(tmp_sub['time2'])
			else:
				subject['seminar']=tmp_sub['class2']
				subject['theory']=tmp_sub['class1']
				day_seminars=self.get_days(tmp_sub['time2'])
				day_theories=self.get_days(tmp_sub['time1'])
			subject['day_theories']=day_theories
			subject['day_seminars']=day_seminars
			self.subjects.append(subject)
	def find_subject(self, name):
		for subject in self.subjects:
			if subject['name']==name:
				return subject
		return None
class DocumentFile:
	def __init__(self, file_contents):
		self.file_contents=file_contents
	def read_file(self):
		document=Document(self.file_contents)
		
		
